from __future__ import annotations

import json
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import resolve_backend_path, settings
from app.models.import_record import ImportRecord
from app.models.taxonomy import Chapter, KnowledgePoint, Subject
from app.models.user import User
from app.schemas.question import QuestionCreate, QuestionOptionIn
from app.services.question_service import create_question, duplicate_candidates
from app.utils.text import normalize_answer, normalize_multi_answer

HEADERS = [
    "题型",
    "题干",
    "选项A",
    "选项B",
    "选项C",
    "选项D",
    "选项E",
    "正确答案",
    "解析",
    "科目",
    "章节",
    "知识点",
    "难度",
    "重要程度",
    "来源",
    "年份",
    "标签",
]

TYPE_MAP = {
    "单选": "single_choice",
    "单选题": "single_choice",
    "single_choice": "single_choice",
    "多选": "multiple_choice",
    "多选题": "multiple_choice",
    "multiple_choice": "multiple_choice",
    "判断": "true_false",
    "判断题": "true_false",
    "true_false": "true_false",
    "填空": "fill_blank",
    "填空题": "fill_blank",
    "fill_blank": "fill_blank",
    "简答": "short_answer",
    "简答题": "short_answer",
    "short_answer": "short_answer",
}
DIFFICULTY_MAP = {"简单": "easy", "easy": "easy", "中等": "medium", "medium": "medium", "困难": "hard", "hard": "hard"}
IMPORTANCE_MAP = {
    "普通": "normal",
    "normal": "normal",
    "重点": "key",
    "高": "key",
    "key": "key",
    "高频": "frequent",
    "高频考点": "frequent",
    "frequent": "frequent",
}
SOURCE_MAP = {
    "自建": "self_built",
    "期末真题": "final_exam",
    "考研": "postgraduate_exam",
    "执医": "physician_exam",
    "规培": "residency_training",
    "其他": "other",
}


def ensure_dirs() -> None:
    resolve_backend_path(settings.upload_dir).mkdir(parents=True, exist_ok=True)
    resolve_backend_path("assets").mkdir(parents=True, exist_ok=True)


def generate_template(path: str | Path = "assets/question_import_template.xlsx") -> Path:
    ensure_dirs()
    path = Path(path)
    if not path.is_absolute():
        path = resolve_backend_path(str(path))
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "题目导入模板"
    ws.append(HEADERS)
    ws.append(
        [
            "单选",
            "能特异性识别抗原肽-MHC 复合物的细胞是？",
            "B细胞",
            "T细胞",
            "红细胞",
            "血小板",
            "",
            "B",
            "T细胞通过TCR识别抗原肽-MHC复合物。",
            "医学免疫学",
            "T 细胞免疫",
            "TCR识别",
            "简单",
            "高频",
            "自建",
            2026,
            "T细胞,抗原识别,MHC",
        ]
    )
    ws.append(
        [
            "多选",
            "下列属于免疫球蛋白功能的是？",
            "特异性结合抗原",
            "激活补体",
            "通过胎盘",
            "携带氧气",
            "",
            "ABC",
            "免疫球蛋白可结合抗原、激活补体，IgG 可通过胎盘。",
            "医学免疫学",
            "免疫球蛋白",
            "抗体功能",
            "中等",
            "重点",
            "自建",
            2026,
            "抗体,补体",
        ]
    )
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max_length + 4, 45)
    wb.save(path)
    return path


def parse_excel_to_rows(path: str | Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = []
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in ws[1]]
    for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(value is None or str(value).strip() == "" for value in row):
            continue
        raw = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        rows.append({"row_number": index, "raw": raw})
    return rows


def preview_import(db: Session, file_path: Path, file_name: str, importer: User) -> ImportRecord:
    rows = parse_excel_to_rows(file_path)
    preview_rows = []
    valid_count = 0
    failed_count = 0
    duplicate_count = 0
    for row in rows:
        result = validate_import_row(db, row["raw"], row["row_number"])
        if result["errors"]:
            failed_count += 1
        else:
            valid_count += 1
        if result["duplicate_candidates"]:
            duplicate_count += 1
        preview_rows.append(result)
    record = ImportRecord(
        importer_id=importer.id,
        file_name=file_name,
        stored_file_path=str(file_path),
        status="preview",
        total_count=len(rows),
        success_count=0,
        failed_count=failed_count,
        skipped_count=0,
        duplicate_count=duplicate_count,
        summary_json=json.dumps(
            {
                "rows": preview_rows,
                "valid_count": valid_count,
                "failed_count": failed_count,
                "duplicate_count": duplicate_count,
            },
            ensure_ascii=False,
        ),
    )
    db.add(record)
    db.flush()
    return record


def validate_import_row(db: Session, raw: dict, row_number: int) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    q_type = TYPE_MAP.get(str(raw.get("题型") or "").strip(), "")
    stem = str(raw.get("题干") or "").strip()
    answer = str(raw.get("正确答案") or "").strip()
    analysis = str(raw.get("解析") or "").strip()
    subject_name = str(raw.get("科目") or "").strip()
    chapter_name = str(raw.get("章节") or "").strip()
    kp_name = str(raw.get("知识点") or "").strip()
    difficulty = DIFFICULTY_MAP.get(str(raw.get("难度") or "中等").strip(), "medium")
    importance = IMPORTANCE_MAP.get(str(raw.get("重要程度") or "普通").strip(), "normal")
    source = SOURCE_MAP.get(str(raw.get("来源") or "自建").strip(), str(raw.get("来源") or "self_built").strip())
    year = parse_year(raw.get("年份"))
    tags = [item.strip() for item in str(raw.get("标签") or "").replace("，", ",").split(",") if item.strip()]

    if not q_type:
        errors.append("题型不能为空或格式不支持")
    if not stem:
        errors.append("题干不能为空")
    if not answer:
        errors.append("正确答案不能为空")
    if not analysis:
        errors.append("解析不能为空")
    if not subject_name:
        errors.append("科目不能为空")
    if not chapter_name:
        errors.append("章节不能为空")
    if not kp_name:
        errors.append("知识点不能为空")

    options = []
    for key in ["A", "B", "C", "D", "E"]:
        value = raw.get(f"选项{key}")
        if value is not None and str(value).strip():
            options.append({"option_key": key, "content": str(value).strip(), "sort_order": len(options) + 1})
    if q_type in {"single_choice", "multiple_choice"} and len(options) < 2:
        errors.append("选择题至少需要 2 个选项")
    option_keys = {option["option_key"] for option in options}
    if q_type == "single_choice" and normalize_answer(answer) not in option_keys:
        errors.append("单选题答案必须是已有选项中的一个")
    if q_type == "multiple_choice":
        answers = normalize_multi_answer(answer).split(",") if answer else []
        if not answers or any(item not in option_keys for item in answers):
            errors.append("多选题答案必须在已有选项范围内")
    if q_type == "true_false" and not options:
        options = [
            {"option_key": "A", "content": "正确", "sort_order": 1},
            {"option_key": "B", "content": "错误", "sort_order": 2},
        ]

    subject = db.scalar(select(Subject).where(Subject.name == subject_name, Subject.deleted_at.is_(None)))
    chapter = None
    kp = None
    if not subject:
        warnings.append("科目不存在，可在确认导入时自动创建")
    else:
        chapter = db.scalar(
            select(Chapter).where(
                Chapter.subject_id == subject.id,
                Chapter.name == chapter_name,
                Chapter.deleted_at.is_(None),
            )
        )
        if not chapter:
            warnings.append("章节不存在，可在确认导入时自动创建")
        else:
            kp = db.scalar(
                select(KnowledgePoint).where(
                    KnowledgePoint.chapter_id == chapter.id,
                    KnowledgePoint.name == kp_name,
                    KnowledgePoint.deleted_at.is_(None),
                )
            )
            if not kp:
                warnings.append("知识点不存在，可在确认导入时自动创建")

    duplicates = duplicate_candidates(db, stem, threshold=0.88, limit=5) if stem else []
    return {
        "row_number": row_number,
        "valid": not errors,
        "errors": errors,
        "warnings": warnings,
        "duplicate_candidates": duplicates,
        "normalized": {
            "question_type": q_type,
            "stem": stem,
            "options": options,
            "correct_answer": answer,
            "analysis": analysis,
            "subject_name": subject_name,
            "chapter_name": chapter_name,
            "knowledge_point_name": kp_name,
            "difficulty": difficulty,
            "importance": importance,
            "source": source,
            "year": year,
            "tags": tags,
        },
    }


def confirm_import(
    db: Session,
    record: ImportRecord,
    importer: User,
    auto_create_taxonomy: bool = True,
    skip_duplicates: bool = True,
    taxonomy_mappings: list[dict] | None = None,
) -> ImportRecord:
    summary = json.loads(record.summary_json or "{}")
    rows = summary.get("rows", [])
    success_count = 0
    failed_count = 0
    skipped_count = 0
    duplicate_count = 0
    error_rows = []
    taxonomy_mappings = taxonomy_mappings or []
    for row in rows:
        if row.get("errors"):
            failed_count += 1
            error_rows.append(row)
            continue
        if row.get("duplicate_candidates"):
            duplicate_count += 1
            if skip_duplicates:
                skipped_count += 1
                continue
        try:
            normalized = row["normalized"]
            mapping = find_taxonomy_mapping(taxonomy_mappings, row)
            if mapping and mapping.get("skip"):
                skipped_count += 1
                continue
            subject, chapter, kp = get_or_create_taxonomy(
                db,
                normalized["subject_name"],
                normalized["chapter_name"],
                normalized["knowledge_point_name"],
                auto_create_taxonomy,
                mapping,
            )
            payload = QuestionCreate(
                stem=normalized["stem"],
                question_type=normalized["question_type"],
                options=[QuestionOptionIn(**option) for option in normalized["options"]],
                correct_answer=normalized["correct_answer"],
                analysis=normalized["analysis"],
                subject_id=subject.id,
                chapter_id=chapter.id,
                knowledge_point_id=kp.id,
                difficulty=normalized["difficulty"],
                importance=normalized["importance"],
                source=normalized["source"],
                year=normalized["year"],
                tags=normalized["tags"],
                status="published",
            )
            create_question(db, payload, importer)
            success_count += 1
        except Exception as exc:  # noqa: BLE001
            failed_count += 1
            row.setdefault("errors", []).append(str(exc))
            error_rows.append(row)
    record.status = "completed"
    record.success_count = success_count
    record.failed_count = failed_count
    record.skipped_count = skipped_count
    record.duplicate_count = duplicate_count
    if error_rows:
        record.error_report_path = str(write_error_report(record.id, error_rows))
    record.summary_json = json.dumps(
        {
            "rows": rows,
            "success_count": success_count,
            "failed_count": failed_count,
            "skipped_count": skipped_count,
            "duplicate_count": duplicate_count,
        },
        ensure_ascii=False,
    )
    return record


def get_or_create_taxonomy(
    db: Session,
    subject_name: str,
    chapter_name: str,
    kp_name: str,
    auto_create: bool,
    mapping: dict | None = None,
) -> tuple[Subject, Chapter, KnowledgePoint]:
    if mapping and (
        mapping.get("target_subject_id")
        or mapping.get("target_chapter_id")
        or mapping.get("target_knowledge_point_id")
    ):
        subject = db.get(Subject, mapping.get("target_subject_id")) if mapping.get("target_subject_id") else None
        chapter = db.get(Chapter, mapping.get("target_chapter_id")) if mapping.get("target_chapter_id") else None
        kp = (
            db.get(KnowledgePoint, mapping.get("target_knowledge_point_id"))
            if mapping.get("target_knowledge_point_id")
            else None
        )
        if not subject and chapter:
            subject = chapter.subject
        if not chapter and kp:
            chapter = kp.chapter
        if not subject and kp:
            subject = kp.subject
        if not subject or subject.deleted_at:
            raise ValueError("手动匹配的科目不存在")
        if not chapter or chapter.deleted_at or chapter.subject_id != subject.id:
            raise ValueError("手动匹配的章节不存在或不属于目标科目")
        if not kp or kp.deleted_at or kp.chapter_id != chapter.id:
            raise ValueError("手动匹配的知识点不存在或不属于目标章节")
        return subject, chapter, kp

    subject = db.scalar(select(Subject).where(Subject.name == subject_name, Subject.deleted_at.is_(None)))
    if not subject:
        if not auto_create:
            raise ValueError(f"科目不存在：{subject_name}")
        subject = Subject(name=subject_name, code=make_code(subject_name), is_enabled=True)
        db.add(subject)
        db.flush()
    chapter = db.scalar(
        select(Chapter).where(
            Chapter.subject_id == subject.id,
            Chapter.name == chapter_name,
            Chapter.deleted_at.is_(None),
        )
    )
    if not chapter:
        if not auto_create:
            raise ValueError(f"章节不存在：{chapter_name}")
        chapter = Chapter(subject_id=subject.id, name=chapter_name, code=make_code(chapter_name), is_enabled=True)
        db.add(chapter)
        db.flush()
    kp = db.scalar(
        select(KnowledgePoint).where(
            KnowledgePoint.chapter_id == chapter.id,
            KnowledgePoint.name == kp_name,
            KnowledgePoint.deleted_at.is_(None),
        )
    )
    if not kp:
        if not auto_create:
            raise ValueError(f"知识点不存在：{kp_name}")
        kp = KnowledgePoint(
            subject_id=subject.id,
            chapter_id=chapter.id,
            name=kp_name,
            label=kp_name,
            importance_level="medium",
            is_enabled=True,
        )
        db.add(kp)
        db.flush()
    return subject, chapter, kp


def find_taxonomy_mapping(mappings: list[dict], row: dict) -> dict | None:
    normalized = row.get("normalized", {})
    row_number = row.get("row_number")
    for mapping in mappings:
        if mapping.get("row_number") and mapping.get("row_number") == row_number:
            return mapping
        if (
            mapping.get("subject_name") == normalized.get("subject_name")
            and mapping.get("chapter_name") == normalized.get("chapter_name")
            and mapping.get("knowledge_point_name") == normalized.get("knowledge_point_name")
        ):
            return mapping
    return None


def write_error_report(import_record_id: int, rows: list[dict]) -> Path:
    ensure_dirs()
    path = resolve_backend_path(settings.upload_dir) / f"import_errors_{import_record_id}_{uuid4().hex[:8]}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "错误报告"
    ws.append(["行号", "错误", "警告", *HEADERS])
    for row in rows:
        normalized = row.get("normalized", {})
        values = [
            row.get("row_number"),
            "；".join(row.get("errors", [])),
            "；".join(row.get("warnings", [])),
            normalized.get("question_type"),
            normalized.get("stem"),
            *[
                next((opt.get("content") for opt in normalized.get("options", []) if opt.get("option_key") == key), "")
                for key in ["A", "B", "C", "D", "E"]
            ],
            normalized.get("correct_answer"),
            normalized.get("analysis"),
            normalized.get("subject_name"),
            normalized.get("chapter_name"),
            normalized.get("knowledge_point_name"),
            normalized.get("difficulty"),
            normalized.get("importance"),
            normalized.get("source"),
            normalized.get("year"),
            ",".join(normalized.get("tags", [])),
        ]
        ws.append(values)
    wb.save(path)
    return path


def parse_year(value) -> int | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return int(float(value))
    except ValueError:
        return None


def make_code(name: str) -> str:
    return f"auto-{abs(hash(name)) % 1000000}"
